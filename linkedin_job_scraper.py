#!/usr/bin/env python3
"""
LinkedIn Visa-Sponsorship Job Scraper
=====================================

Scrapes software-engineering jobs across multiple countries using LinkedIn's
PUBLIC guest endpoint (no login required -> won't get your account flagged),
tags each job with a visa-sponsorship signal, and writes a clean Excel workbook
AND an interactive HTML report, one sheet/tab per country.

WHY THIS APPROACH
-----------------
LinkedIn aggressively blocks logged-in automation. But the "guest" jobs API
(the same one that powers the public, logged-out job search) is open and
returns parseable job cards. We use only that. Be polite: keep the delays.

SETUP
-----
    pip install requests beautifulsoup4 openpyxl

RUN
---
    python linkedin_job_scraper.py

Edit the CONFIG block below to control countries, roles, and the time window.
"""

import time
import random
import datetime as dt
from urllib.parse import urlencode

import requests
from bs4 import BeautifulSoup
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

# =============================================================================
# CONFIG  --  edit this block, nothing else needs changing
# =============================================================================

# How fresh should jobs be?
# Use any number of hours: "1h", "6h", "12h", "24h", "48h", etc.
# Or "week" for the last 7 days.
TIME_WINDOW = "24h"

# Countries to search. The "location" string is what LinkedIn matches against.
COUNTRIES = [
    {"name": "Germany",     "location": "Germany"},
    {"name": "Spain",       "location": "Spain"},
    {"name": "London",      "location": "London, England, United Kingdom"},
    {"name": "Netherlands", "location": "Netherlands"},
]

# Role variants. Each one is searched separately so you don't miss postings
# that only match one phrasing. Add more if you like.
ROLE_KEYWORDS = [
    "software engineer",
    "backend engineer",
    "full stack software engineer",
]

# For every role above, we ALSO run a "<role> visa sponsorship" search so we
# catch postings that explicitly advertise sponsorship. Keep True for max recall.
ALSO_SEARCH_VISA_KEYWORD = True

# Fetch each job's full description to detect visa/sponsorship language and tag
# the row STRONG / WEAK / NONE. Slower (one extra request per unique job) but
# this is what stops you wasting time on non-sponsoring roles. Recommended True.
CHECK_DESCRIPTIONS = True

# Output files
OUTPUT_FILE = "linkedin_jobs.xlsx"
HTML_OUTPUT_FILE = "linkedin_jobs.html"

# Politeness / anti-rate-limit. Don't crank these down or LinkedIn will 429 you.
PAGE_DELAY = (2.5, 4.5)      # seconds between paginated requests (randomised)
DESC_DELAY = (1.0, 2.0)      # seconds between description fetches
MAX_PAGES_PER_SEARCH = 10    # ~25 jobs/page, so up to ~250 per search
REQUEST_TIMEOUT = 20

# =============================================================================
# Internals
# =============================================================================

def _parse_time_window(tw):
    """Return (f_TPR value, human label) for any 'Nh' or 'week' string."""
    tw = tw.strip().lower()
    if tw == "week":
        return "r604800", "last 7 days"
    if tw.endswith("h"):
        try:
            hours = float(tw[:-1])
            seconds = int(hours * 3600)
            label = f"last {tw[:-1]} hour{'s' if hours != 1 else ''}"
            return f"r{seconds}", label
        except ValueError:
            pass
    raise SystemExit(
        f"Invalid TIME_WINDOW '{tw}'. Use 'Nh' (e.g. '1h', '6h', '24h') or 'week'."
    )

SEARCH_URL = (
    "https://www.linkedin.com/jobs-guest/jobs/api/seeMoreJobPostings/search"
)
JOB_DETAIL_URL = "https://www.linkedin.com/jobs-guest/jobs/api/jobPosting/{job_id}"

HEADERS = {
    "User-Agent": (
        "Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 "
        "(KHTML, like Gecko) Chrome/125.0.0.0 Safari/537.36"
    ),
    "Accept": "text/html,application/xhtml+xml,application/xml;q=0.9,*/*;q=0.8",
    "Accept-Language": "en-US,en;q=0.9",
}

# Visa / relocation signal vocabulary, weighted.
STRONG_VISA_TERMS = [
    "visa sponsorship", "sponsor your visa", "we sponsor", "sponsorship available",
    "will sponsor", "visa sponsor", "sponsored visa", "relocation package",
    "work permit sponsorship", "blue card", "skilled worker visa", "tier 2",
    "visa support", "sponsor work visa", "sponsorship provided",
]
WEAK_VISA_TERMS = [
    "relocation", "relocate", "visa", "work permit", "sponsorship",
    "international candidates", "willing to relocate", "permit",
]

# Words that, appearing just before a visa term, flip its meaning
# ("we do NOT offer visa sponsorship" must not be tagged STRONG).
NEGATORS = [
    "no ", "not ", "without", "cannot", "can't", "unable",
    "don't", "doesn't", "won't", "neither", "nor ", "never",
    "unfortunately",
]

# The subset of terms that are specifically about *sponsorship*. When one of
# these is negated ("no visa sponsorship", "we do not sponsor"), it's an
# explicit rejection -> NO-SPONSOR, rather than mere silence -> NONE.
# Generic words like "visa"/"relocation" are deliberately excluded so a phrase
# like "no visa required" doesn't get mislabelled as a rejection.
SPONSOR_TERMS = {
    "visa sponsorship", "sponsor your visa", "we sponsor", "sponsorship available",
    "will sponsor", "visa sponsor", "sponsored visa", "work permit sponsorship",
    "visa support", "sponsor work visa", "sponsorship provided", "sponsorship",
    "blue card", "skilled worker visa", "tier 2",
}

session = requests.Session()
session.headers.update(HEADERS)


def _sleep(rng):
    time.sleep(random.uniform(*rng))


def build_search_url(keywords, location, tpr, start):
    params = {
        "keywords": keywords,
        "location": location,
        "f_TPR": tpr,
        "start": start,
    }
    return f"{SEARCH_URL}?{urlencode(params)}"


def parse_cards(html):
    """Pull job cards out of a guest-API HTML chunk."""
    soup = BeautifulSoup(html, "html.parser")
    cards = soup.select("li")
    jobs = []
    for li in cards:
        base = li.find("div", class_=lambda c: c and "base-card" in c)
        if not base:
            continue

        # Job id from the entity urn
        urn = base.get("data-entity-urn", "")
        job_id = urn.split(":")[-1] if urn else None

        title_el = li.find(class_="base-search-card__title")
        company_el = li.find(class_="base-search-card__subtitle")
        location_el = li.find(class_="job-search-card__location")
        link_el = li.find("a", class_="base-card__full-link")
        date_el = li.find("time")

        if not (title_el and link_el):
            continue

        url = link_el.get("href", "").split("?")[0]
        posted = ""
        if date_el:
            posted = date_el.get("datetime", "") or date_el.get_text(strip=True)

        jobs.append({
            "job_id": job_id,
            "title": title_el.get_text(strip=True),
            "company": company_el.get_text(strip=True) if company_el else "",
            "location": location_el.get_text(strip=True) if location_el else "",
            "posted": posted,
            "url": url,
        })
    return jobs


def fetch_description(job_id):
    """Fetch a single job's full text for visa-signal scanning."""
    if not job_id:
        return ""
    try:
        r = session.get(
            JOB_DETAIL_URL.format(job_id=job_id), timeout=REQUEST_TIMEOUT
        )
        if r.status_code != 200:
            return ""
        soup = BeautifulSoup(r.text, "html.parser")
        desc = soup.find(class_="show-more-less-html__markup") or soup.find(
            class_="description__text"
        )
        return desc.get_text(" ", strip=True).lower() if desc else ""
    except requests.RequestException:
        return ""


def _negated(t, term):
    """True only if EVERY occurrence of `term` is preceded by a negator."""
    idx = t.find(term)
    if idx == -1:
        return False
    while idx != -1:
        window = t[max(0, idx - 35):idx]
        if not any(neg in window for neg in NEGATORS):
            return False          # a clean, non-negated mention exists
        idx = t.find(term, idx + 1)
    return True


def visa_signal(text):
    """Classify a description's visa stance. Returns one of:

      STRONG      - sponsorship clearly offered
      WEAK        - visa / relocation mentioned in passing
      NONE        - description read fine, but no visa language at all
      NO-SPONSOR  - explicitly states it will NOT sponsor
      UNKNOWN     - description couldn't be read (nothing to judge)

    Negation-aware: a sponsorship term preceded by a negator
    ('no visa sponsorship', 'we do not sponsor') is treated as an explicit
    rejection (NO-SPONSOR), not as a positive signal. A clean positive
    mention anywhere wins over a negated one.
    """
    if not text:
        return "UNKNOWN"
    t = text.lower()
    found_negated_sponsor = False

    for term in STRONG_VISA_TERMS:
        if term in t:
            if _negated(t, term):
                if term in SPONSOR_TERMS:
                    found_negated_sponsor = True
            else:
                return "STRONG"

    for term in WEAK_VISA_TERMS:
        if term in t:
            if _negated(t, term):
                if term in SPONSOR_TERMS:
                    found_negated_sponsor = True
            else:
                return "WEAK"

    return "NO-SPONSOR" if found_negated_sponsor else "NONE"


def search_one(keywords, location, tpr):
    """Paginate a single keyword+location search; return list of job dicts."""
    found = []
    start = 0
    for _ in range(MAX_PAGES_PER_SEARCH):
        url = build_search_url(keywords, location, tpr, start)
        try:
            r = session.get(url, timeout=REQUEST_TIMEOUT)
        except requests.RequestException as e:
            print(f"      ! request error: {e}")
            break

        if r.status_code == 429:
            print("      ! rate-limited (429) — backing off 30s")
            time.sleep(30)
            continue
        if r.status_code != 200 or not r.text.strip():
            break

        page_jobs = parse_cards(r.text)
        if not page_jobs:
            break

        found.extend(page_jobs)
        start += len(page_jobs)
        _sleep(PAGE_DELAY)

    return found


def scrape():
    tpr, window_label = _parse_time_window(TIME_WINDOW)
    print(f"Scraping LinkedIn jobs ({window_label})\n" + "=" * 50)

    # country -> {job_id: job_dict}
    results = {c["name"]: {} for c in COUNTRIES}

    for country in COUNTRIES:
        name, loc = country["name"], country["location"]
        print(f"\n[{name}]")

        # Build the full keyword list for this country
        queries = list(ROLE_KEYWORDS)
        if ALSO_SEARCH_VISA_KEYWORD:
            queries += [f"{role} visa sponsorship" for role in ROLE_KEYWORDS]

        for kw in queries:
            print(f"  searching: '{kw}'")
            jobs = search_one(kw, loc, tpr)
            new = 0
            for j in jobs:
                jid = j["job_id"] or j["url"]
                if jid not in results[name]:
                    results[name][jid] = j
                    new += 1
            print(f"      +{new} new (total {len(results[name])})")

    # Optional: enrich with visa signal from descriptions
    if CHECK_DESCRIPTIONS:
        print("\nChecking descriptions for visa signals...")
        for name, jobs in results.items():
            todo = [j for j in jobs.values() if j["job_id"]]
            print(f"  [{name}] {len(todo)} jobs")
            for i, job in enumerate(todo, 1):
                txt = fetch_description(job["job_id"])
                job["visa"] = visa_signal(txt)
                if i % 10 == 0:
                    print(f"      {i}/{len(todo)}")
                _sleep(DESC_DELAY)
    else:
        for jobs in results.values():
            for job in jobs.values():
                job["visa"] = "UNCHECKED"

    return results, window_label


# =============================================================================
# Excel output
# =============================================================================

VISA_ORDER = {"STRONG": 0, "WEAK": 1, "UNKNOWN": 2, "UNCHECKED": 2,
              "NONE": 3, "NO-SPONSOR": 4}
VISA_FILL = {
    "STRONG":     PatternFill("solid", fgColor="C6EFCE"),   # green
    "WEAK":       PatternFill("solid", fgColor="FFEB9C"),   # amber
    "NONE":       PatternFill("solid", fgColor="FFC7CE"),   # light red (no signal)
    "NO-SPONSOR": PatternFill("solid", fgColor="BFBFBF"),   # grey (explicitly ruled out)
}
HEADER_FILL = PatternFill("solid", fgColor="1F2937")
HEADER_FONT = Font(color="FFFFFF", bold=True)
COLS = ["Visa", "Title", "Company", "Location", "Posted", "Job URL"]


def write_excel(results, window_label):
    wb = Workbook()
    wb.remove(wb.active)

    total = 0
    for country in COUNTRIES:
        name = country["name"]
        jobs = list(results[name].values())
        # sort: strongest visa signal first, then company
        jobs.sort(key=lambda j: (VISA_ORDER.get(j.get("visa"), 2),
                                 j.get("company", "").lower()))

        ws = wb.create_sheet(title=name[:31])

        # header
        for ci, col in enumerate(COLS, 1):
            cell = ws.cell(row=1, column=ci, value=col)
            cell.fill = HEADER_FILL
            cell.font = HEADER_FONT
            cell.alignment = Alignment(vertical="center")

        for ri, job in enumerate(jobs, 2):
            visa = job.get("visa", "UNKNOWN")
            ws.cell(row=ri, column=1, value=visa)
            ws.cell(row=ri, column=2, value=job["title"])
            ws.cell(row=ri, column=3, value=job["company"])
            ws.cell(row=ri, column=4, value=job["location"])
            ws.cell(row=ri, column=5, value=job["posted"])
            link = ws.cell(row=ri, column=6, value=job["url"])
            link.hyperlink = job["url"]
            link.font = Font(color="2563EB", underline="single")

            fill = VISA_FILL.get(visa)
            if fill:
                ws.cell(row=ri, column=1).fill = fill

        # column widths
        widths = [10, 42, 26, 30, 22, 60]
        for ci, w in enumerate(widths, 1):
            ws.column_dimensions[get_column_letter(ci)].width = w
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = f"A1:{get_column_letter(len(COLS))}{len(jobs)+1}"

        total += len(jobs)
        print(f"  [{name}] wrote {len(jobs)} jobs")

    # summary sheet up front
    summary = wb.create_sheet(title="Summary", index=0)
    summary["A1"] = "LinkedIn Job Scrape"
    summary["A1"].font = Font(bold=True, size=14)
    summary["A2"] = f"Window: {window_label}"
    summary["A3"] = f"Generated: {dt.datetime.now():%Y-%m-%d %H:%M}"
    summary["A5"] = "Country"
    summary["B5"] = "Jobs"
    summary["C5"] = "Strong visa"
    for c in ("A5", "B5", "C5"):
        summary[c].font = Font(bold=True)
    for i, country in enumerate(COUNTRIES, 6):
        name = country["name"]
        jobs = results[name].values()
        strong = sum(1 for j in jobs if j.get("visa") == "STRONG")
        summary.cell(row=i, column=1, value=name)
        summary.cell(row=i, column=2, value=len(jobs))
        summary.cell(row=i, column=3, value=strong)
    summary.column_dimensions["A"].width = 16
    summary.column_dimensions["B"].width = 10
    summary.column_dimensions["C"].width = 12

    wb.save(OUTPUT_FILE)
    print(f"\nDone. {total} jobs saved to {OUTPUT_FILE}")


# =============================================================================
# HTML output
# =============================================================================

_VISA_BADGE = {
    "STRONG":     ('<span class="badge badge-strong">&#10003; STRONG</span>', 0),
    "WEAK":       ('<span class="badge badge-weak">&#9679; WEAK</span>',      1),
    "UNKNOWN":    ('<span class="badge badge-unknown">? UNKNOWN</span>',      2),
    "UNCHECKED":  ('<span class="badge badge-unknown">&#8212; UNCHECKED</span>', 2),
    "NONE":       ('<span class="badge badge-none">&#10007; NONE</span>',     3),
    "NO-SPONSOR": ('<span class="badge badge-nosponsor">&#8856; NO-SPONSOR</span>', 4),
}

_HTML_TEMPLATE = """\
<!DOCTYPE html>
<html lang="en">
<head>
<meta charset="UTF-8"/>
<meta name="viewport" content="width=device-width, initial-scale=1.0"/>
<title>LinkedIn Jobs &mdash; {window_label}</title>
<style>
  *, *::before, *::after {{ box-sizing: border-box; margin: 0; padding: 0; }}
  :root {{
    --bg: #0f172a; --surface: #1e293b; --surface2: #293548;
    --border: #334155; --text: #e2e8f0; --muted: #94a3b8;
    --accent: #3b82f6; --accent-hover: #2563eb;
    --strong-bg: #14532d; --strong-fg: #86efac;
    --weak-bg:   #713f12; --weak-fg:   #fde68a;
    --none-bg:   #7f1d1d; --none-fg:   #fca5a5;
    --nosp-bg:   #111827; --nosp-fg:   #6b7280;
    --unk-bg:    #1e293b; --unk-fg:    #94a3b8;
    --radius: 8px; --font: "Inter", "Segoe UI", system-ui, sans-serif;
  }}
  body {{ background: var(--bg); color: var(--text); font-family: var(--font);
          font-size: 14px; min-height: 100vh; }}

  /* ---- header ---- */
  header {{
    background: var(--surface); border-bottom: 1px solid var(--border);
    padding: 18px 28px; display: flex; align-items: center; gap: 16px;
    position: sticky; top: 0; z-index: 100; flex-wrap: wrap;
  }}
  header h1 {{ font-size: 20px; font-weight: 700; color: #fff; flex: 1; white-space: nowrap; }}
  header .meta {{ font-size: 12px; color: var(--muted); }}
  .search-wrap {{ display: flex; align-items: center; gap: 8px; }}
  #search {{
    background: var(--bg); border: 1px solid var(--border); border-radius: var(--radius);
    color: var(--text); padding: 7px 12px; font-size: 13px; width: 240px; outline: none;
    transition: border-color .15s;
  }}
  #search:focus {{ border-color: var(--accent); }}
  #search::placeholder {{ color: var(--muted); }}
  .filter-btns {{ display: flex; gap: 6px; flex-wrap: wrap; }}
  .filter-btn {{
    padding: 5px 12px; border-radius: 99px; border: 1px solid var(--border);
    background: transparent; color: var(--muted); cursor: pointer; font-size: 12px;
    transition: all .15s;
  }}
  .filter-btn.active, .filter-btn:hover {{ background: var(--accent); border-color: var(--accent); color: #fff; }}
  .filter-btn.active {{ font-weight: 600; }}

  /* ---- tabs ---- */
  .tabs {{ display: flex; gap: 2px; padding: 16px 28px 0; border-bottom: 1px solid var(--border); }}
  .tab {{
    padding: 10px 20px; border-radius: var(--radius) var(--radius) 0 0;
    cursor: pointer; color: var(--muted); font-size: 13px; font-weight: 500;
    border: 1px solid transparent; border-bottom: none; transition: all .15s;
    background: transparent; user-select: none;
  }}
  .tab:hover {{ color: var(--text); background: var(--surface2); }}
  .tab.active {{
    background: var(--surface); color: #fff; border-color: var(--border);
    border-bottom: 1px solid var(--surface);  margin-bottom: -1px;
  }}
  .tab-count {{
    display: inline-block; background: var(--surface2); border-radius: 99px;
    padding: 1px 7px; font-size: 11px; margin-left: 6px; color: var(--muted);
  }}
  .tab.active .tab-count {{ background: var(--accent); color: #fff; }}

  /* ---- panels ---- */
  .panel {{ display: none; padding: 24px 28px; }}
  .panel.active {{ display: block; }}

  /* ---- stats bar ---- */
  .stats {{ display: flex; gap: 12px; margin-bottom: 16px; flex-wrap: wrap; }}
  .stat {{
    background: var(--surface); border: 1px solid var(--border); border-radius: var(--radius);
    padding: 10px 16px; font-size: 12px; color: var(--muted);
  }}
  .stat strong {{ display: block; font-size: 20px; font-weight: 700; color: var(--text); }}

  /* ---- table ---- */
  .table-wrap {{ overflow-x: auto; border-radius: var(--radius); border: 1px solid var(--border); }}
  table {{ width: 100%; border-collapse: collapse; }}
  thead tr {{ background: var(--surface); }}
  th {{
    padding: 11px 14px; text-align: left; font-size: 11px; font-weight: 600;
    text-transform: uppercase; letter-spacing: .05em; color: var(--muted);
    border-bottom: 1px solid var(--border); white-space: nowrap;
    cursor: pointer; user-select: none;
  }}
  th:hover {{ color: var(--text); }}
  th .sort-icon {{ opacity: .4; margin-left: 4px; }}
  th.sorted .sort-icon {{ opacity: 1; color: var(--accent); }}
  tbody tr {{ border-bottom: 1px solid var(--border); transition: background .1s; }}
  tbody tr:last-child {{ border-bottom: none; }}
  tbody tr:hover {{ background: var(--surface2); }}
  tbody tr.hidden {{ display: none; }}
  td {{ padding: 11px 14px; vertical-align: middle; }}
  td.title {{ font-weight: 500; color: #fff; max-width: 340px; }}
  td.company {{ color: var(--muted); }}
  td.location {{ color: var(--muted); font-size: 12px; }}
  td.posted {{ color: var(--muted); font-size: 12px; white-space: nowrap; }}
  .job-link {{
    display: inline-flex; align-items: center; gap: 5px;
    color: var(--accent); text-decoration: none; font-size: 12px;
    font-weight: 500; white-space: nowrap; padding: 4px 10px;
    border: 1px solid var(--accent); border-radius: var(--radius);
    transition: all .15s;
  }}
  .job-link:hover {{ background: var(--accent); color: #fff; }}
  .job-link svg {{ width: 12px; height: 12px; flex-shrink: 0; }}

  /* ---- badges ---- */
  .badge {{
    display: inline-block; padding: 3px 9px; border-radius: 99px;
    font-size: 11px; font-weight: 600; letter-spacing: .03em; white-space: nowrap;
  }}
  .badge-strong {{ background: var(--strong-bg); color: var(--strong-fg); }}
  .badge-weak   {{ background: var(--weak-bg);   color: var(--weak-fg);   }}
  .badge-none   {{ background: var(--none-bg);   color: var(--none-fg);   }}
  .badge-nosponsor {{ background: var(--nosp-bg); color: var(--nosp-fg); border: 1px solid #374151; text-decoration: line-through; }}
  .badge-unknown{{ background: var(--unk-bg);    color: var(--unk-fg); border: 1px solid var(--border); }}

  /* ---- empty state ---- */
  .empty {{ text-align: center; padding: 60px 20px; color: var(--muted); }}
  .empty svg {{ width: 40px; height: 40px; margin-bottom: 12px; opacity: .3; }}
</style>
</head>
<body>

<header>
  <div>
    <h1>&#128188; LinkedIn Job Results</h1>
    <div class="meta">Window: {window_label} &nbsp;&bull;&nbsp; Generated: {generated}</div>
  </div>
  <div class="search-wrap">
    <input id="search" type="text" placeholder="Search title, company, location&hellip;" oninput="applyFilters()"/>
  </div>
  <div class="filter-btns" id="visaFilters">
    <button class="filter-btn active" data-visa="ALL"    onclick="setVisa(this)">All</button>
    <button class="filter-btn"        data-visa="STRONG" onclick="setVisa(this)">&#10003; Strong</button>
    <button class="filter-btn"        data-visa="WEAK"   onclick="setVisa(this)">&#9679; Weak</button>
    <button class="filter-btn"        data-visa="UNKNOWN" onclick="setVisa(this)">? Unknown</button>
    <button class="filter-btn"        data-visa="NONE"   onclick="setVisa(this)">&#10007; None</button>
    <button class="filter-btn"        data-visa="NO-SPONSOR" onclick="setVisa(this)">&#8856; No-sponsor</button>
  </div>
</header>

<div class="tabs" id="tabs">
{tab_html}
</div>

{panels_html}

<script>
var activeVisa = "ALL";
var activeTab  = 0;

function switchTab(idx) {{
  document.querySelectorAll(".tab").forEach(function(t,i)   {{ t.classList.toggle("active", i===idx); }});
  document.querySelectorAll(".panel").forEach(function(p,i) {{ p.classList.toggle("active", i===idx); }});
  activeTab = idx;
  applyFilters();
}}

function setVisa(btn) {{
  document.querySelectorAll(".filter-btn").forEach(function(b) {{ b.classList.remove("active"); }});
  btn.classList.add("active");
  activeVisa = btn.dataset.visa;
  applyFilters();
}}

function applyFilters() {{
  var q = document.getElementById("search").value.toLowerCase();
  var panel = document.querySelectorAll(".panel")[activeTab];
  if (!panel) return;
  var rows = panel.querySelectorAll("tbody tr");
  rows.forEach(function(row) {{
    var visa   = row.dataset.visa || "";
    var text   = row.dataset.text || "";
    var visaOk = activeVisa === "ALL" || visa === activeVisa;
    var textOk = !q || text.includes(q);
    row.classList.toggle("hidden", !(visaOk && textOk));
  }});
}}

// simple column sort
function sortTable(thEl, colIdx) {{
  var table  = thEl.closest("table");
  var tbody  = table.querySelector("tbody");
  var rows   = Array.from(tbody.querySelectorAll("tr"));
  var asc    = thEl.dataset.asc !== "1";
  thEl.dataset.asc = asc ? "1" : "0";
  table.querySelectorAll("th").forEach(function(t) {{ t.classList.remove("sorted"); var ic = t.querySelector(".sort-icon"); if (ic) ic.textContent = " \\u2195"; }});
  thEl.classList.add("sorted");
  thEl.querySelector(".sort-icon").textContent = asc ? " \\u2191" : " \\u2193";
  rows.sort(function(a, b) {{
    var av = a.querySelectorAll("td")[colIdx] ? a.querySelectorAll("td")[colIdx].textContent.trim() : "";
    var bv = b.querySelectorAll("td")[colIdx] ? b.querySelectorAll("td")[colIdx].textContent.trim() : "";
    return asc ? av.localeCompare(bv) : bv.localeCompare(av);
  }});
  rows.forEach(function(r) {{ tbody.appendChild(r); }});
}}
</script>
</body>
</html>
"""

_PANEL_TEMPLATE = """\
<div class="panel{active_class}" id="panel-{idx}">
  <div class="stats">
    <div class="stat"><strong>{total}</strong>Total jobs</div>
    <div class="stat"><strong style="color:var(--strong-fg)">{n_strong}</strong>Strong visa</div>
    <div class="stat"><strong style="color:var(--weak-fg)">{n_weak}</strong>Weak signal</div>
    <div class="stat"><strong style="color:var(--none-fg)">{n_none}</strong>No signal</div>
    <div class="stat"><strong style="color:var(--nosp-fg)">{n_nosp}</strong>Ruled out</div>
  </div>
  <div class="table-wrap">
    <table>
      <thead>
        <tr>
          <th onclick="sortTable(this,0)">Visa<span class="sort-icon">&#8597;</span></th>
          <th onclick="sortTable(this,1)">Title<span class="sort-icon">&#8597;</span></th>
          <th onclick="sortTable(this,2)">Company<span class="sort-icon">&#8597;</span></th>
          <th onclick="sortTable(this,3)">Location<span class="sort-icon">&#8597;</span></th>
          <th onclick="sortTable(this,4)">Posted<span class="sort-icon">&#8597;</span></th>
          <th>Link</th>
        </tr>
      </thead>
      <tbody>
{rows_html}
      </tbody>
    </table>
  </div>
</div>
"""

_LINK_SVG = (
    '<svg viewBox="0 0 24 24" fill="none" stroke="currentColor" stroke-width="2.5">'
    '<path d="M18 13v6a2 2 0 0 1-2 2H5a2 2 0 0 1-2-2V8a2 2 0 0 1 2-2h6"/>'
    '<polyline points="15 3 21 3 21 9"/><line x1="10" y1="14" x2="21" y2="3"/>'
    '</svg>'
)

def _esc(s):
    """Minimal HTML escaping."""
    return (str(s)
            .replace("&", "&amp;")
            .replace("<", "&lt;")
            .replace(">", "&gt;")
            .replace('"', "&quot;"))


def write_html(results, window_label):
    tab_parts   = []
    panel_parts = []

    for idx, country in enumerate(COUNTRIES):
        name = country["name"]
        jobs = list(results[name].values())
        jobs.sort(key=lambda j: (VISA_ORDER.get(j.get("visa"), 2),
                                  j.get("company", "").lower()))

        n_strong = sum(1 for j in jobs if j.get("visa") == "STRONG")
        n_weak   = sum(1 for j in jobs if j.get("visa") == "WEAK")
        n_none   = sum(1 for j in jobs if j.get("visa") == "NONE")
        n_nosp   = sum(1 for j in jobs if j.get("visa") == "NO-SPONSOR")

        # Tab button
        tab_parts.append(
            f'<div class="tab{" active" if idx == 0 else ""}" '
            f'onclick="switchTab({idx})">'
            f'{_esc(name)}'
            f'<span class="tab-count">{len(jobs)}</span>'
            f'</div>'
        )

        # Table rows
        rows = []
        for job in jobs:
            visa  = job.get("visa", "UNKNOWN")
            badge, _ = _VISA_BADGE.get(visa, _VISA_BADGE["UNKNOWN"])
            search_text = _esc(
                f"{job.get('title','')} {job.get('company','')} {job.get('location','')}".lower()
            )
            url = _esc(job.get("url", ""))
            rows.append(
                f'        <tr data-visa="{_esc(visa)}" data-text="{search_text}">\n'
                f'          <td>{badge}</td>\n'
                f'          <td class="title">{_esc(job.get("title",""))}</td>\n'
                f'          <td class="company">{_esc(job.get("company",""))}</td>\n'
                f'          <td class="location">{_esc(job.get("location",""))}</td>\n'
                f'          <td class="posted">{_esc(job.get("posted",""))}</td>\n'
                f'          <td><a class="job-link" href="{url}" target="_blank" rel="noopener">'
                f'{_LINK_SVG} View</a></td>\n'
                f'        </tr>'
            )

        panel_parts.append(_PANEL_TEMPLATE.format(
            active_class=" active" if idx == 0 else "",
            idx=idx,
            total=len(jobs),
            n_strong=n_strong,
            n_weak=n_weak,
            n_none=n_none,
            n_nosp=n_nosp,
            rows_html="\n".join(rows) if rows else
                      '        <tr><td colspan="6" class="empty">No jobs found</td></tr>',
        ))

    html = _HTML_TEMPLATE.format(
        window_label=_esc(window_label),
        generated=dt.datetime.now().strftime("%Y-%m-%d %H:%M"),
        tab_html="\n".join(tab_parts),
        panels_html="\n".join(panel_parts),
    )

    with open(HTML_OUTPUT_FILE, "w", encoding="utf-8") as fh:
        fh.write(html)
    print(f"HTML report saved to {HTML_OUTPUT_FILE}")


if __name__ == "__main__":
    results, window_label = scrape()
    try:
        write_excel(results, window_label)
    except PermissionError:
        print(f"Warning: could not save {OUTPUT_FILE} — close it in Excel and re-run if you need it.")
    write_html(results, window_label)